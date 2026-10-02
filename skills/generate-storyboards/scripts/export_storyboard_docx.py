#!/usr/bin/env python3
"""
export_storyboard_docx.py - Generates a styled, executive-ready DOCX and HTML deliverable
from a storyboard concept and its visual graphic.

Usage:
    python3 export_storyboard_docx.py <image_path> <title> [output_prefix]
"""

import sys
import base64
from pathlib import Path

def export_storyboard_html(image_path: Path, title: str, output_prefix: Path):
    """Generates a clean HTML deliverable suitable for copying into Google Docs or sharing."""
    html_file = output_prefix.with_suffix(".html")
    
    img_tag = ""
    if image_path.exists():
        try:
            mime = "image/jpeg" if image_path.suffix.lower() in [".jpg", ".jpeg"] else "image/png"
            b64_data = base64.b64encode(image_path.read_bytes()).decode("utf-8")
            img_src = f"data:{mime};base64,{b64_data}"
            img_tag = f'<div style="text-align: center; margin: 24px 0;"><img src="{img_src}" alt="{title}" style="max-width: 100%; height: auto; border-radius: 8px; box-shadow: 0 4px 12px rgba(0,0,0,0.1);"/></div>'
        except Exception:
            img_tag = f'<div style="text-align: center; margin: 24px 0;"><img src="{image_path.name}" alt="{title}" style="max-width: 100%; height: auto;"/></div>'
            
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    max-width: 960px;
    margin: 40px auto;
    padding: 0 20px;
    background-color: #ffffff;
  }}
  h1 {{
    color: #334155;
    text-align: center;
    font-size: 28px;
    margin-bottom: 24px;
  }}
</style>
</head>
<body>
  <h1>{title}</h1>
  {img_tag}
</body>
</html>
"""
    html_file.write_text(html_content, encoding="utf-8")
    print(f"Exported HTML: {html_file}")

def export_storyboard_docx(image_path: Path, title: str, output_prefix: Path):
    """Generates an executive-styled DOCX document."""
    try:
        import docx
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("Notice: python-docx not installed. Skipping DOCX generation (install with `pip install python-docx`).")
        return

    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    COLOR_PRIMARY = RGBColor(51, 65, 85)     # #334155 Slate
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = 'Arial'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    p.paragraph_format.space_after = Pt(12)
    
    if image_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(8)
        run_img = p_img.add_run()
        run_img.add_picture(str(image_path), width=Inches(6.8))
        
    docx_file = output_prefix.with_suffix(".docx")
    doc.save(str(docx_file))
    print(f"Exported DOCX: {docx_file}")

def export_storyboard(image_path: Path, title: str, output_prefix: Path):
    export_storyboard_docx(image_path, title, output_prefix)
    export_storyboard_html(image_path, title, output_prefix)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python3 export_storyboard_docx.py <image_path> <title> [output_prefix]")
        sys.exit(1)
    img = Path(sys.argv[1])
    ttl = sys.argv[2]
    out = Path(sys.argv[3]) if len(sys.argv) > 3 else img.with_suffix("")
    export_storyboard(img, ttl, out)
